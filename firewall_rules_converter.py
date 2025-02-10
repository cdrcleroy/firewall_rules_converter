"""
🔥 Firewall Rules Converter v1.0 🔥
Développé par Cedric Leroy - 2025
Contact : cdrc.leroy@gmail.com
"""

import tkinter as tk
from tkinter import filedialog
import xml.etree.ElementTree as ET
import re
import time
from tqdm import tqdm
import sys
import openpyxl
from openpyxl.styles import Font
from openpyxl.worksheet.filters import AutoFilter
from openpyxl.styles import PatternFill
from datetime import datetime

WATERMARK = "© 2025 - Développé par Cedric Leroy - firewall_rules_converter"

def get_watermark():
    return WATERMARK

def select_file(file_type):
    """Ouvre un explorateur de fichiers pour sélectionner un fichier."""
    file_path = filedialog.askopenfilename(title=f"Sélectionner le fichier {file_type}", filetypes=[("Fichiers", "*.*")])
    return file_path if file_path else None

def select_output_file():
    """Ouvre un explorateur de fichiers pour choisir l'emplacement et le nom du fichier XLSX de sortie."""
    default_filename = "firewall_" + datetime.today().strftime('%Y-%d-%m') + ".xlsx"
    file_path = filedialog.asksaveasfilename(defaultextension=".xlsx", filetypes=[("Excel files", "*.xlsx")], title="Enregistrer le fichier XLSX", initialfile=default_filename)
    return file_path if file_path else default_filename

def get_text(element, default=""):
    """Récupère le texte d'un élément XML avec une valeur par défaut."""
    return element.text if element is not None else default

def get_all_texts(elements):
    """Retourne une liste des textes de plusieurs éléments XML."""
    return [el.text for el in elements if el.text]

def clean_list(value_list):
    """Supprime les doublons et enlève 'N/A' si d'autres valeurs sont présentes."""
    cleaned_list = list(dict.fromkeys(value_list))
    if "N/A" in cleaned_list and len(cleaned_list) > 1:
        cleaned_list.remove("N/A")
    return cleaned_list

def clean_ip_list(ip_list, ip_loaded):
    """Si 'any' est présent et qu'un fichier d'IP est chargé, ne conserve que 'any'. Si aucun fichier IP n'est chargé, retourne 'N/A'."""
    if not ip_loaded:
        return ["N/A"]
    cleaned_list = list(dict.fromkeys(ip_list))
    return ["any"] if "any" in cleaned_list else cleaned_list

def parse_ip_hosts(xml_file):
    """Analyse le fichier XML des adresses IP et retourne un dictionnaire associant les noms aux IPs."""
    print("Chargement du fichier d'objets réseaux...")
    time.sleep(0.5)
    tree = ET.parse(xml_file)
    root = tree.getroot()
    return {
        get_text(host.find("Name")): get_text(host.find("IPAddress"), f"{get_text(host.find('StartIPAddress'))} - {get_text(host.find('EndIPAddress'))}")
        for host in root.findall(".//IPHost")
    }

def parse_services(xml_file):
    """Analyse le fichier XML des services et retourne un dictionnaire associant les services à leurs détails."""
    print("Chargement du fichier des services...")
    time.sleep(0.5)
    tree = ET.parse(xml_file)
    root = tree.getroot()
    return {
        get_text(service.find("Name")): [
            (get_text(detail.find("Protocol"), get_text(detail.find("ProtocolName"), "N/A")),
             get_text(detail.find("DestinationPort"), "N/A"))
            for detail in service.findall(".//ServiceDetail")
        ]
        for service in root.findall(".//Services")
    }

def parse_positions(position_file):
    """Analyse le fichier texte des positions et retourne un dictionnaire associant la position à l'ID."""
    print("Chargement du fichier des id...")
    time.sleep(0.5)
    position_data = {}
    with open(position_file, "r", encoding="utf-8") as file:
        for line in file:
            match = re.search(r"--id (\d+) --position (\d+)", line)
            if match:
                position_data[int(match.group(2))] = match.group(1)
    return position_data

def parse_firewall_rules(xml_file, ip_data, service_data, position_data, ip_loaded):
    """Analyse le fichier XML et extrait les règles de firewall en les croisant avec les IPs, services et positions."""
    print("Traitement des règles de firewall...")
    tree = ET.parse(xml_file)
    root = tree.getroot()
    rules = []
    firewall_rules = list(root.findall(".//FirewallRule"))
    
    for position, rule in enumerate(tqdm(firewall_rules, desc="Analyse des règles", unit="règle"), start=1):
        source_networks = get_all_texts(rule.findall(".//SourceNetworks/Network"))
        destination_networks = get_all_texts(rule.findall(".//DestinationNetworks/Network"))
        services = get_all_texts(rule.findall(".//Services/Service"))
        
        source_ips = clean_ip_list([ip_data.get(net, "any") for net in source_networks], ip_loaded)
        destination_ips = clean_ip_list([ip_data.get(net, "any") for net in destination_networks], ip_loaded)
        service_details = [service_data.get(srv, [("N/A", "N/A")]) for srv in services]
        
        flat_service_details = [item for sublist in service_details for item in sublist]
        destination_ports = clean_list([d[1] for d in flat_service_details])
        protocols = clean_list([d[0] for d in flat_service_details])
        
        rule_data = {
            "ID": position_data.get(position, "N/A"),
            "Name": get_text(rule.find("Name")),
            "Description": get_text(rule.find("Description")),
            "IPFamily": get_text(rule.find("IPFamily")),
            "Status": get_text(rule.find("Status")),
            "PolicyType": get_text(rule.find("PolicyType")),
            "Action": get_text(rule.find(".//Action")),
            "LogTraffic": get_text(rule.find(".//LogTraffic")),
            "SourceZones": "; ".join(get_all_texts(rule.findall(".//SourceZones/Zone"))),
            "DestinationZones": "; ".join(get_all_texts(rule.findall(".//DestinationZones/Zone"))),
            "Schedule": get_text(rule.find(".//Schedule")),
            "Exclusions": "",
            "SourceNetworks": "; ".join(source_networks),
            "Source Network IP": "; ".join(source_ips),
            "DestinationNetworks": "; ".join(destination_networks),
            "Destination Network IP": "; ".join(destination_ips),
            "Services": "; ".join(services),
            "Destination Ports": "; ".join(destination_ports),
            "Protocols": "; ".join(protocols)
        }
        rules.append(rule_data)
    
    print("Traitement terminé avec succès.")
    return rules

def save_to_excel(rules, output_file):
    if not rules:
        print("Aucune règle à sauvegarder.")
        return
    
    fieldnames = ["ID", "Name", "Description", "IPFamily", "Status", "PolicyType", "Action", "LogTraffic", "SourceZones", "DestinationZones", "Schedule", "Exclusions", "SourceNetworks", "Source Network IP", "DestinationNetworks", "Destination Network IP", "Services", "Destination Ports", "Protocols"]
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Firewall Rules"
    ws.append(fieldnames)
    
    for col_num, column_title in enumerate(fieldnames, 1):
        cell = ws.cell(row=1, column=col_num, value=column_title)
        cell.font = Font(bold=True)
        ws.column_dimensions[openpyxl.utils.get_column_letter(col_num)].width = max(len(column_title) + 2, 15)
    
    fill_grey = PatternFill(start_color="DDDDDD", end_color="DDDDDD", fill_type="solid")
    for row_idx, rule in enumerate(rules, start=2):
        row_values = [rule.get(field, "N/A") for field in fieldnames]
        ws.append(row_values)
        if row_idx % 2 == 0:  # Appliquer la couleur de fond toutes les deux lignes
            for col_idx in range(1, len(fieldnames) + 1):
                ws.cell(row=row_idx, column=col_idx).fill = fill_grey
        
    ws.auto_filter.ref = ws.dimensions    
        
    try:
        wb.save(output_file)
        print(f"✅ Fichier Excel enregistré : {output_file}")
    except Exception as e:
        print(f"❌ Erreur : {e}")
        
def main():
    print("=== Menu de Chargement des Fichiers Firewall ===")
    print("Choisissez un fichier de règles firewall :")
    
    firewall_rules_file = select_file("des règles firewall")
    if not firewall_rules_file:
        print("Aucun fichier XML sélectionné. Fin du programme.")
        sys.exit(0)
    
    load_positions = input("Voulez-vous charger un fichier des id des règles ? (O/N) : ").strip().lower()
    position_data = parse_positions(select_file("des positions")) if load_positions == 'o' else {}
    
    load_ip = input("Voulez-vous charger un fichier d'objets réseaux ? (O/N) : ").strip().lower()
    ip_data = parse_ip_hosts(select_file("des IP")) if load_ip == 'o' else {}
    ip_loaded = load_ip == 'o'
    
    load_services = input("Voulez-vous charger un fichier de services ? (O/N) : ").strip().lower()
    service_data = parse_services(select_file("des Services")) if load_services == 'o' else {}
    
    output_file = select_output_file()
    
    rules = parse_firewall_rules(firewall_rules_file, ip_data, service_data, position_data, ip_loaded)
    save_to_excel(rules, output_file)
    
if __name__ == "__main__":
    tk.Tk().withdraw()
    main()
